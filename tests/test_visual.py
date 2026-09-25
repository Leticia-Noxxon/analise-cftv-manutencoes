"""Teste visual (spec §45, caso 6) no site compilado: manutenção em célula vermelha mantém o fundo
vermelho e mostra o marcador azul. Também verifica que o site carrega sem erros de JavaScript.

Requer `site/dist` (cd site && npm run build) e Playwright com Chrome/Chromium; caso contrário é pulado."""
import sys
from pathlib import Path

import pytest

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / 'scripts'))
pw = pytest.importorskip('playwright.sync_api')

if not (RAIZ / 'site' / 'dist' / 'index.html').exists():
    pytest.skip('site/dist não compilado', allow_module_level=True)

VERMELHO = 'rgb(214, 59, 59)'
AZUL = 'rgb(27, 42, 107)'
BRANCO = 'rgb(255, 255, 255)'


@pytest.fixture(scope='module')
def pagina():
    from servidor_teste import iniciar
    srv, url = iniciar()
    with pw.sync_playwright() as p:
        try:
            b = p.chromium.launch(channel='chrome')
        except Exception:
            try:
                b = p.chromium.launch()
            except Exception as e:  # pragma: no cover
                pytest.skip(f'navegador indisponível: {e}')
        pg = b.new_page(viewport={'width': 1500, 'height': 900}, accept_downloads=True)
        erros = []
        pg.on('pageerror', lambda e: erros.append(str(e)))
        pg.goto(url)
        pg.wait_for_selector('body[data-pronto="1"]')
        pg.erros = erros
        yield pg
        b.close()
    srv.shutdown()


def _caso(pg, status):
    """Primeiro (prefixo, coluna) com manutenção cuja célula tem o status pedido."""
    return pg.evaluate("""(st) => { for (const p of window.__D.P) for (const j of Object.keys(p.evByCol))
        if (p.g[j] === st) return [p.p, +j]; return null; }""", status)


def _celula(pg, prefixo, j):
    pg.fill('#f-busca', '')
    pg.fill('#f-busca', str(prefixo))
    pg.wait_for_function('n => document.querySelectorAll(".mx-row").length === 1 && document.querySelector(".pf b").textContent == n', arg=str(prefixo))
    return pg.locator(f'.c[data-i="0"][data-j="{j}"]')


def test_6_visual_manutencao_em_celula_vermelha(pagina):
    caso = _caso(pagina, 'R')
    assert caso, 'não há manutenção em dia vermelho nos dados'
    cel = _celula(pagina, *caso)
    assert 'sR' in cel.get_attribute('class')
    assert cel.evaluate('e => getComputedStyle(e).backgroundColor') == VERMELHO
    dot = cel.locator('.dot')
    assert dot.count() == 1
    assert dot.evaluate('e => getComputedStyle(e).backgroundColor') == AZUL
    # o marcador não cobre a célula inteira (a cor de fundo continua visível)
    cb, db = cel.bounding_box(), dot.bounding_box()
    assert db['width'] * db['height'] < 0.5 * cb['width'] * cb['height']


def test_6b_visual_manutencao_em_dia_sem_dados_fica_branca_com_marcador(pagina):
    caso = _caso(pagina, ' ')
    assert caso
    cel = _celula(pagina, *caso)
    assert cel.evaluate('e => getComputedStyle(e).backgroundColor') == BRANCO
    assert cel.locator('.dot').count() == 1


def test_site_abas_sem_erros_e_todos_os_veiculos(pagina):
    pagina.fill('#f-busca', '')
    for aba in ['manutencoes', 'efetividade', 'recorrencias', 'auditoria', 'como-ler', 'matriz']:
        pagina.click(f'[data-tab="{aba}"]')
        pagina.wait_for_timeout(150)
    pagina.click('[data-veic="todos"]')
    pagina.wait_for_timeout(200)
    assert '5.430 veículo' in pagina.inner_text('#f-count')
    # virtualização: só as linhas visíveis ficam no DOM
    assert pagina.locator('.mx-row').count() < 100
    assert pagina.erros == []


def test_exportacao_excel_tem_as_5_planilhas(pagina, tmp_path):
    import openpyxl
    pagina.fill('#f-busca', '')
    pagina.click('[data-veic="com"]')
    with pagina.expect_download() as d:
        pagina.click('#btn-exp')
    arq = tmp_path / 'exp.xlsx'
    d.value.save_as(arq)
    wb = openpyxl.load_workbook(arq, read_only=True)
    assert wb.sheetnames == ['Resumo', 'Histórico por prefixo', 'Manutenções', 'Resultado das manutenções', 'Recorrências']
    assert wb['Manutenções'].max_row == 325  # 324 formulários + cabeçalho
