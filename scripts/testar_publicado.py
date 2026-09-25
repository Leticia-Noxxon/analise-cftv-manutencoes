"""Fase 17 – teste do dashboard publicado (URL pública) em navegador headless com perfil limpo.

Uso:  python scripts/testar_publicado.py [URL]
Salva capturas em docs/screenshots/publicado_*.png e imprime o resultado de cada verificação."""
import sys
import tempfile
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

URL = sys.argv[1] if len(sys.argv) > 1 else 'https://leticia-noxxon.github.io/analise-cftv-manutencoes/'
OUT = Path(__file__).resolve().parents[1] / 'docs' / 'screenshots'
res = []


def ok(nome, cond, extra=''):
    res.append((nome, bool(cond), extra))
    print(('OK   ' if cond else 'FALHA'), nome, extra)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p, tempfile.TemporaryDirectory() as perfil:
        try:
            ctx = p.chromium.launch_persistent_context(perfil, channel='chrome', viewport={'width': 1600, 'height': 950}, accept_downloads=True)
        except Exception:
            ctx = p.chromium.launch_persistent_context(perfil, viewport={'width': 1600, 'height': 950}, accept_downloads=True)
        pg = ctx.new_page()
        falhas_rede, erros_js, respostas = [], [], {}
        pg.on('response', lambda r: (respostas.__setitem__(r.url, r.status), r.status >= 400 and falhas_rede.append(f'{r.status} {r.url}')))
        pg.on('requestfailed', lambda r: falhas_rede.append(f'FALHOU {r.url}'))
        pg.on('pageerror', lambda e: erros_js.append(str(e)))
        pg.goto(URL, wait_until='networkidle')
        pg.wait_for_selector('body[data-pronto="1"]', timeout=60000)
        dados = {u.split('/')[-1]: s for u, s in respostas.items() if '/data/' in u}
        ok('JSON de dados carregados (200)', dados.get('cftv.json') == 200 and dados.get('manutencoes.json') == 200 and dados.get('meta.json') == 200, str(dados))
        ok('assets JS/CSS carregados', any('/assets/' in u and u.endswith('.js') and s == 200 for u, s in respostas.items()) and any('/assets/' in u and u.endswith('.css') and s == 200 for u, s in respostas.items()))
        n = pg.locator('.mx-row').count()
        ok('matriz renderizada (Com manutenção)', n > 10 and '306 veículo' in pg.inner_text('#f-count'), f'{n} linhas visíveis; {pg.inner_text("#f-count")}')
        pg.screenshot(path=OUT / 'publicado_01_matriz.png')

        pg.fill('#f-busca', '32704')
        pg.wait_for_timeout(500)
        if pg.locator('.mx-row').count() == 0:  # 32704 não tem manutenção: o painel oferece "Todos os veículos"
            ok('busca 32704 na lista "Com manutenção" mostra aviso', pg.locator('#ver-todos').count() == 1, pg.inner_text('.mx-empty').replace('\n', ' '))
            pg.screenshot(path=OUT / 'publicado_02a_busca_sem_manutencao.png')
            pg.click('#ver-todos')
        pg.wait_for_function('() => document.querySelectorAll(".mx-row").length === 1')
        ok('busca 32704', pg.inner_text('.pf b') == '32704', pg.inner_text('#f-count'))
        j = pg.evaluate("() => { const p = window.__D.byPrefixo[32704]; return p.g.split('').findIndex(g => g !== ' '); }")
        cel = pg.locator(f'.c[data-i="0"][data-j="{j}"]')
        cel.hover()
        pg.wait_for_selector('#tooltip:not([hidden])')
        tt = pg.inner_text('#tooltip')
        ok('tooltip com câmeras', 'Prefixo 32704' in tt and 'câm' in tt, tt.replace('\n', ' | ')[:140])
        pg.screenshot(path=OUT / 'publicado_02_tooltip_busca_32704.png')
        cel.click()
        pg.wait_for_selector('#drawer:not([hidden])')
        ok('painel de detalhe da célula', pg.locator('#drawer .tbl tbody tr').count() > 0 and 'DETALHE DO DIA' in pg.inner_text('#drawer'))
        pg.screenshot(path=OUT / 'publicado_03_detalhe_celula.png')
        pg.keyboard.press('Escape')
        pg.locator('.pf').first.click()
        pg.wait_for_selector('#drawer:not([hidden])')
        ok('linha do tempo 32704 (sem manutenção)', 'não recebeu manutenção' in pg.inner_text('#drawer'))
        pg.keyboard.press('Escape')
        # painel de manutenção: exemplo com antes → depois → recorrência
        pg.fill('#f-busca', '31005')
        pg.wait_for_function('() => document.querySelectorAll(".mx-row").length === 1 && document.querySelector(".pf b").textContent === "31005"')
        dots = pg.locator('.mx-row .dot').count()
        if dots:
            pg.locator('.mx-row .dot').first.click()
            pg.wait_for_selector('#drawer:not([hidden])')
            t = pg.inner_text('#drawer')
            ok('painel de manutenção', 'Resultado' in t and 'OBSERVAÇÃO COMPLETA' in t.upper(), pg.inner_text('#drawer .res-line .badge'))
            pg.screenshot(path=OUT / 'publicado_04_painel_manutencao.png')
            pg.keyboard.press('Escape')
        else:
            ok('painel de manutenção', False, 'sem bolinha')
        pg.locator('.pf').first.click()
        pg.wait_for_selector('#drawer:not([hidden])')
        t = pg.inner_text('#drawer')
        ok('linha do tempo do prefixo', 'LINHA DO TEMPO' in t and 'Eventos do período' in t)
        pg.screenshot(path=OUT / 'publicado_05_linha_do_tempo.png')
        pg.keyboard.press('Escape')

        with pg.expect_download() as d:
            pg.click('#btn-exp')
        arq = Path(perfil) / 'exp.xlsx'
        d.value.save_as(arq)
        wb = openpyxl.load_workbook(arq, read_only=True)
        ok('exportação Excel', len(wb.sheetnames) == 5 and wb['Manutenções'].max_row >= 2, f'{wb.sheetnames}; prefixo filtrado → {wb["Manutenções"].max_row - 1} formulário(s)')

        pg.fill('#f-busca', '')
        for aba, nome in [('manutencoes', 'Manutenções'), ('efetividade', 'Taxa de resolução'), ('recorrencias', 'Recorrências na mesma câmera'),
                          ('auditoria', 'Como cada texto de câmera foi interpretado'), ('como-ler', 'Como ler este painel')]:
            pg.click(f'[data-tab="{aba}"]')
            pg.wait_for_timeout(400)
            ok(f'aba {aba}', nome in pg.inner_text('#conteudo'))
            if aba == 'efetividade':
                ok('taxa de resolução 36,4%', '36,4%' in pg.inner_text('#conteudo'))
                pg.screenshot(path=OUT / 'publicado_06_efetividade.png')
        pg.click('[data-tab="matriz"]')
        pg.click('[data-veic="todos"]')
        pg.wait_for_timeout(400)
        ok('todos os veículos', '5.430 veículo' in pg.inner_text('#f-count'), pg.inner_text('#f-count'))
        pg.locator('.mx-scroll').evaluate('e => e.scrollTop = 60000')
        pg.wait_for_timeout(300)
        ok('rolagem com 5.430 linhas', pg.locator('.mx-row').count() > 10)
        pg.screenshot(path=OUT / 'publicado_07_todos_veiculos.png')
        ok('sem 404/falhas de rede', not falhas_rede, '; '.join(falhas_rede[:5]))
        ok('sem erros de JavaScript', not erros_js, '; '.join(erros_js[:3]))
        ctx.close()
    falhas = [r for r in res if not r[1]]
    print(f'\n{len(res) - len(falhas)}/{len(res)} verificações OK')
    sys.exit(1 if falhas else 0)


if __name__ == '__main__':
    main()
