"""Leitura integral do arquivo Revisão_CFTV (todas as abas, todas as linhas, texto completo)."""
import os, glob, unicodedata, json
import openpyxl, pandas as pd
RAW = os.path.join(os.path.dirname(__file__), '..', '..', 'data', 'raw')
f = [x for x in glob.glob(os.path.join(RAW, '*.xlsx')) if unicodedata.normalize('NFC', os.path.basename(x)).startswith('Revisão_CFTV')][0]
wb = openpyxl.load_workbook(f, data_only=False)
for ws in wb.worksheets:
    print('ABA', ws.title, ws.dimensions, 'merged', ws.merged_cells.ranges, 'pivots', len(getattr(ws,'_pivots',[])), 'tables', list(ws.tables.keys()))
    formulas = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=')]
    print('  fórmulas:', len(formulas), formulas[:10])
ws = wb['Planilha1']
for r in ws.iter_rows(values_only=True):
    if any(v is not None for v in r[3:]): print('  P1', r)
print('P1 col B distinct:', pd.Series([r[1] for r in ws.iter_rows(min_row=2, values_only=True)]).value_counts(dropna=False).to_dict())
wb2 = openpyxl.load_workbook(f, data_only=True)
ws = wb2['Sheet1']
rows = list(ws.iter_rows(values_only=True))
hdr = list(rows[0]); data = rows[1:]
print('linhas de dados', len(data), 'linhas totalmente vazias', sum(all(v is None for v in r) for r in data))
df = pd.DataFrame(data, columns=hdr)
df.insert(0, 'linha_excel', range(2, len(data)+2))
df.to_pickle(os.path.join(os.path.dirname(__file__), '..', '..', 'data', 'processed', '_manut_raw.pkl'))
for c in hdr:
    s = df[c]
    nn = s.notna() & (s.astype(str).str.strip() != '')
    print('\n###', repr(c), 'preenchidos', nn.sum(), f'({nn.mean():.0%})', 'tipos', s[nn].map(lambda x: type(x).__name__).value_counts().to_dict(), 'distintos', s[nn].nunique())
    vc = s[nn].astype(str).value_counts()
    if len(vc) <= 40:
        for k, v in vc.items(): print(f'   {v:4d} | {k!r}')
    else:
        for k, v in vc.head(12).items(): print(f'   {v:4d} | {k!r}')
    # itens separados por \n
    if s[nn].astype(str).str.contains('\n').any() and c != 'Observações':
        items = s[nn].astype(str).str.split('\n').explode().str.strip()
        print('   ITENS (split por quebra de linha):')
        for k, v in items.value_counts().items(): print(f'      {v:4d} | {k!r}')
