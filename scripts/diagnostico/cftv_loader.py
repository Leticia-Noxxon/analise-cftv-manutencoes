"""Leitura dos arquivos 'Relatório CFTV' (somente leitura de data/raw).
Identifica automaticamente a aba de registros (detalhe) e ignora abas de tabela dinâmica/resumo."""
import glob, os, re, unicodedata
import openpyxl
import pandas as pd

RAW = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..', 'data', 'raw'))

def norm_header(h):
    if h is None:
        return None
    s = unicodedata.normalize('NFC', str(h)).strip()
    s = re.sub(r'\s+', ' ', s)
    return s

def list_cftv_files():
    files = [f for f in glob.glob(os.path.join(RAW, '*.xlsx'))
             if unicodedata.normalize('NFC', os.path.basename(f)).startswith('Relatório CFTV')]
    return sorted(files, key=lambda f: date_from_filename(f) or pd.Timestamp.max)

def date_from_filename(f):
    m = re.search(r'(\d{2})\.(\d{2})\.(\d{4})', os.path.basename(f))
    return pd.Timestamp(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None

def classify_sheet(ws):
    """Retorna (tipo, motivo). tipo in {'registros','ignorada'}"""
    pivots = len(getattr(ws, '_pivots', []))
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    hdr = [norm_header(h) for h in first]
    has_prefixo = 'Prefixo' in hdr
    has_cam = any(h and re.match(r'^C[âa]mera \d+$', h) for h in hdr)
    if pivots:
        return 'ignorada', f'contém {pivots} tabela(s) dinâmica(s) (resumo "Contagem de Prefixo" por Empresa × Status); não é base de registros'
    if has_prefixo and has_cam:
        return 'registros', 'cabeçalho na linha 1 com Prefixo + colunas Câmera NN; uma linha por veículo'
    return 'ignorada', 'não possui cabeçalho de registros (Prefixo + Câmera NN)'

def load_file(f):
    wb = openpyxl.load_workbook(f, data_only=True)
    info = {'arquivo': os.path.basename(f), 'abas': [], 'aba_usada': None, 'abas_ignoradas': []}
    df = None
    for ws in wb.worksheets:
        tipo, motivo = classify_sheet(ws)
        info['abas'].append(ws.title)
        if tipo == 'registros' and df is None:
            rows = list(ws.iter_rows(values_only=True))
            raw_hdr = list(rows[0])
            info['colunas_originais'] = raw_hdr
            hdr = [norm_header(h) for h in raw_hdr]
            data = [r for r in rows[1:] if any(v is not None for v in r)]
            df = pd.DataFrame(data, columns=hdr)
            df.insert(0, 'linha_excel', [i + 2 for i, r in enumerate(rows[1:]) if any(v is not None for v in r)])
            info['aba_usada'] = ws.title
            info['motivo_aba'] = motivo
            info['linhas_vazias_ignoradas'] = len(rows) - 1 - len(data)
        else:
            info['abas_ignoradas'].append((ws.title, motivo))
    df.insert(0, 'arquivo_origem', info['arquivo'])
    df.insert(1, 'data_arquivo', date_from_filename(f))
    return df, info

def load_all():
    dfs, infos = [], []
    for f in list_cftv_files():
        df, info = load_file(f)
        dfs.append(df); infos.append(info)
    return dfs, infos
