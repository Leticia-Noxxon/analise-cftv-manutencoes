"""Confere o total geral da aba 'Tabela' (tabela dinâmica) com o nº de registros da aba de detalhe."""
import os, sys, glob
sys.path.insert(0, os.path.dirname(__file__))
import openpyxl
from cftv_loader import list_cftv_files
for f in list_cftv_files():
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb['Tabela']
    rows = list(ws.iter_rows(values_only=True))
    tot = [r for r in rows if r and r[0] == 'Total Geral']
    det = wb.worksheets[1]
    n = sum(1 for r in det.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r))
    print(os.path.basename(f), 'Total Geral pivot =', tot[0][len([x for x in tot[0] if x is not None]) - 1] if tot else None, '| detalhe =', n)
