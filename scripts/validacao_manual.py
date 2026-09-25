"""Fase 13 – validação por amostragem: planilha original (openpyxl, leitura independente) × bases
processadas (data/processed) × JSON do site × o que o painel exibe (Playwright). Gera docs/VALIDACAO.md.

Uso:  python scripts/validacao_manual.py   (requer site/dist compilado)
"""
import datetime as dt
import json
import random
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import pandas as pd
from playwright.sync_api import sync_playwright

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / 'scripts'))
from servidor_teste import iniciar  # noqa: E402

RAW = RAIZ / 'data' / 'raw'
PROC = RAIZ / 'data' / 'processed'
SITE = RAIZ / 'site' / 'public' / 'data'
CAMS = [21, 22, 23, 24, 25, 26]
SEMENTE = 20260925


def status_independente(textos):
    """Regra da especificação reimplementada de forma independente do pipeline."""
    existentes = [t.strip() for t in textos if t is not None and str(t).strip() not in ('', '-')]
    if not existentes:
        return 'S'
    if any(t.upper().startswith('OFFLINE') for t in existentes):
        return 'R'
    if any('error' in t.lower() for t in existentes):
        return 'L'
    return 'V'


def ler_cftv_bruto(alvos):
    regs = defaultdict(dict)  # prefixo -> data -> dict
    for arq in sorted(RAW.glob('Relatório CFTV - *.xlsx')):
        data = dt.datetime.strptime(re.search(r'(\d\d\.\d\d\.\d{4})', arq.name).group(1), '%d.%m.%Y').date()
        wb = openpyxl.load_workbook(arq, read_only=True, data_only=True)
        ws = next(w for w in wb.worksheets if w.title.startswith('Relatório'))
        linhas = ws.iter_rows(values_only=True)
        cab = [str(c).strip() if c is not None else '' for c in next(linhas)]
        ip = cab.index('Prefixo')
        ic = {n: cab.index(f'Câmera {n}') for n in CAMS}
        ie = cab.index('Empresa')
        for num, row in enumerate(linhas, start=2):
            try:
                p = int(row[ip])
            except (TypeError, ValueError):
                continue
            if p in alvos:
                textos = [row[ic[n]] for n in CAMS]
                regs[p][data] = {'linha': num, 'arquivo': arq.name, 'empresa': row[ie], 'textos': textos, 'status': status_independente(textos)}
        wb.close()
    return regs


def ler_manut_bruto():
    arq = next(RAW.glob('Revisão_CFTV*.xlsx'))
    wb = openpyxl.load_workbook(arq, read_only=True, data_only=True)
    ws = wb['Sheet1']
    linhas = ws.iter_rows(values_only=True)
    cab = [str(c).strip() if c is not None else '' for c in next(linhas)]
    out = defaultdict(list)
    for num, row in enumerate(linhas, start=2):
        d = dict(zip(cab, row))
        try:
            p = int(d['Prefixo'])
        except (TypeError, ValueError):
            continue
        out[p].append({'linha': num, 'obs': d.get('Observações') or '', 'nome': d.get('Nome'), 'data': d.get('Data')})
    wb.close()
    return out


def main():
    cftv = json.load(open(SITE / 'cftv.json', encoding='utf-8'))
    man = json.load(open(SITE / 'manutencoes.json', encoding='utf-8'))
    cols = [c['data'] for c in cftv['colunas']]
    itens = {x['p']: x for x in cftv['prefixos']}
    com_m = sorted({e['prefixo'] for e in man['eventos'] if e['prefixo_no_cftv']})
    sem_m = sorted(set(itens) - {e['prefixo'] for e in man['eventos']})
    rnd = random.Random(SEMENTE)
    amostra = rnd.sample(com_m, 6) + rnd.sample(sem_m, 4)
    # garante casos especiais: um com recorrência, um com mais de um formulário e o prefixo ausente do CFTV
    extra = [next(e['prefixo'] for e in man['eventos'] if e['recorrencia']),
             next(e['prefixo'] for e in man['eventos'] if e['qtd_formularios'] > 1),
             next(e['prefixo'] for e in man['eventos'] if not e['prefixo_no_cftv']),
             max(man['forms'], key=lambda f: len(f.get('observacoes') or ''))['prefixo']]
    for p in extra:
        if p not in amostra:
            amostra.append(p)
    alvos = set(amostra)
    bruto = ler_cftv_bruto(alvos)
    mbruto = ler_manut_bruto()
    proc = pd.read_csv(PROC / 'cftv_consolidado.csv', dtype={'prefixo': int}, usecols=lambda c: c in ('prefixo', 'data', 'status_geral', 'linha_excel', 'arquivo_origem') or c.endswith('_original'))
    proc = proc[proc['prefixo'].isin(alvos)]
    forms_proc = pd.read_csv(PROC / 'manutencoes_tratadas.csv')

    resultados = []
    srv, url = iniciar()
    with sync_playwright() as pw:
        try:
            b = pw.chromium.launch(channel='chrome')
        except Exception:
            b = pw.chromium.launch()
        pg = b.new_page(viewport={'width': 1500, 'height': 950})
        pg.goto(url)
        pg.wait_for_selector('body[data-pronto="1"]')
        pg.click('[data-veic="todos"]')
        for p in amostra:
            r = {'prefixo': p, 'problemas': [], 'com_manut': p in com_m or p in [e['prefixo'] for e in man['eventos']]}
            it = itens[p]
            # 1) bruto × processado × JSON (cor de cada dia)
            dias_ok = 0
            for j, iso in enumerate(cols):
                d = dt.date.fromisoformat(iso)
                esperado = bruto.get(p, {}).get(d, {}).get('status', ' ')
                pj = it['g'][j]
                pp = proc[(proc.prefixo == p) & (proc.data == iso)]
                pc = pp['status_geral'].iloc[0] if len(pp) else ' '
                if not (esperado == pj == pc):
                    r['problemas'].append(f'{iso}: bruto={esperado!r} processado={pc!r} json={pj!r}')
                else:
                    dias_ok += 1
                if len(pp) and esperado != ' ':
                    br = bruto[p][d]
                    if int(pp['linha_excel'].iloc[0]) != br['linha']:
                        r['problemas'].append(f'{iso}: linha Excel difere ({pp["linha_excel"].iloc[0]} × {br["linha"]})')
                    for n, t in zip(CAMS, br['textos']):
                        tp = pp[f'cam{n}_original'].iloc[0]
                        tb = '' if t is None else str(t)
                        tp = '' if pd.isna(tp) else str(tp)
                        if tb.strip() != tp.strip():
                            r['problemas'].append(f'{iso} câm {n}: texto original difere')
            r['dias'] = f'{dias_ok}/{len(cols)}'
            r['registros_brutos'] = len(bruto.get(p, {}))
            r['cores'] = ''.join(it['g'])
            # 2) painel: cores exibidas
            pg.fill('#f-busca', str(p))
            pg.wait_for_function('n => document.querySelector(".pf b") && document.querySelector(".pf b").textContent == n && document.querySelectorAll(".mx-row").length === 1', arg=str(p))
            classes = pg.eval_on_selector_all('.mx-row .c', 'es => es.map(e => (e.className.match(/s[VLRBI]/)||["?"])[0])')
            mapa = {'sV': 'V', 'sL': 'L', 'sR': 'R', 'sB': ' ', 'sI': 'I'}
            painel = ''.join(mapa.get(c, '?') for c in classes)
            painel_esp = it['g'].replace('S', ' ')
            if painel != painel_esp:
                r['problemas'].append(f'painel mostra {painel!r} esperado {painel_esp!r}')
            n_dots = pg.locator('.mx-row .dot').count()
            # 3) detalhe de uma célula com dados: textos por câmera
            jdata = next((j for j, g in enumerate(it['g']) if g in 'RL'), next((j for j, g in enumerate(it['g']) if g not in ' '), None))
            if jdata is not None:
                pg.locator(f'.c[data-i="0"][data-j="{jdata}"]').click()
                pg.wait_for_selector('#drawer:not([hidden])')
                linhas = pg.eval_on_selector_all('#drawer .tbl tbody tr', 'rs => rs.map(r => [...r.cells].map(c => c.innerText.trim()))')
                br = bruto[p][dt.date.fromisoformat(cols[jdata])]
                exist = [(n, t) for n, t in zip(CAMS, br['textos']) if t is not None and str(t).strip() not in ('', '-')]
                if len(linhas) != len(exist):
                    r['problemas'].append(f'painel da célula {cols[jdata]}: {len(linhas)} câmeras exibidas × {len(exist)} existentes na planilha')
                for (n, t), lin in zip(exist, linhas):
                    t = str(t)
                    exp = 'OFFLINE' if t.upper().startswith('OFFLINE') else ('ONLINE COM ERRO' if 'error' in t.lower() else 'ONLINE')
                    if f'câm {n}' not in lin[0] or lin[1] != exp:
                        r['problemas'].append(f'painel da célula {cols[jdata]} câm {n}: exibido {lin[:2]} × planilha {t!r}')
                    for k, nome in zip((2, 3, 4), ('SD', 'Login', 'Gravação')):
                        m = re.search(nome + r':\s*(\w+)', t)
                        if m and lin[k] != m.group(1):
                            r['problemas'].append(f'câm {n} {nome}: exibido {lin[k]} × planilha {m.group(1)}')
                r['celula_conferida'] = cols[jdata]
                pg.keyboard.press('Escape')
            # 4) manutenções: planilha × processado × painel (texto integral da observação)
            evs = [e for e in man['eventos'] if e['prefixo'] == p]
            mb = mbruto.get(p, [])
            r['forms_planilha'] = len(mb)
            r['forms_processados'] = int((forms_proc['prefixo'] == p).sum())
            r['forms_json'] = sum(e['qtd_formularios'] for e in evs)
            if not (r['forms_planilha'] == r['forms_processados'] == r['forms_json']):
                r['problemas'].append(f"formulários: planilha {r['forms_planilha']} × processado {r['forms_processados']} × json {r['forms_json']}")
            r['bolinhas'] = n_dots
            if n_dots != len([e for e in evs if e['data'] in cols]):
                r['problemas'].append(f'{n_dots} bolinhas × {len(evs)} eventos')
            r['resultados'] = []
            obs_painel = []
            for k in range(n_dots):
                pg.locator('.mx-row .dot').nth(k).click()
                pg.wait_for_selector('#drawer:not([hidden])')
                r['resultados'].append(pg.inner_text('#drawer .res-line .badge'))
                obs_painel += pg.eval_on_selector_all('#drawer .obs', 'es => es.map(e => e.textContent)')
                pg.keyboard.press('Escape')
            obs_bruto = [str(x['obs']) for x in mb if x['obs']]
            faltando = [o for o in obs_bruto if o not in obs_painel]
            if faltando:
                r['problemas'].append(f'{len(faltando)} observação(ões) da planilha não aparecem idênticas no painel')
            r['obs'] = f'{len(obs_bruto) - len(faltando)}/{len(obs_bruto)} idênticas (maior: {max((len(o) for o in obs_bruto), default=0)} caracteres)'
            r['resultados_json'] = [e['resultado'] for e in evs]
            if n_dots and r['resultados'] != [e['resultado'] for e in sorted(evs, key=lambda e: e['data']) if e['data'] in cols]:
                r['problemas'].append(f"resultado exibido {r['resultados']} × json {r['resultados_json']}")
            resultados.append(r)
        b.close()
    srv.shutdown()
    escrever(resultados)


def escrever(res):
    ok = sum(1 for r in res if not r['problemas'])
    L = ['# VALIDAÇÃO MANUAL POR AMOSTRAGEM (Fase 13)', '',
         f'Gerado por `scripts/validacao_manual.py` em {dt.datetime.now():%d/%m/%Y %H:%M} (horário de Brasília).', '',
         f'Amostra aleatória com semente fixa ({SEMENTE}): 6 prefixos com manutenção + 4 sem manutenção, acrescida de casos especiais '
         '(recorrência, mais de um formulário no mesmo dia, prefixo ausente do CFTV, observação mais longa do arquivo) quando não sorteados.', '',
         '**O que foi comparado em cada prefixo:**', '',
         '1. **Planilha original** (lida de novo com openpyxl, sem usar o código do pipeline) → cor esperada de cada dia, recalculada com a regra da especificação '
         '(OFFLINE → vermelho; “error” sem OFFLINE → laranja; tudo ONLINE ok → verde; “-” e vazio ignorados; sem linha → branco).',
         '2. **Base processada** `data/processed/cftv_consolidado.csv` → status, linha do Excel e texto original de cada câmera.',
         '3. **JSON do site** `site/public/data/cftv.json` → cor de cada dia.',
         '4. **Painel** (navegador headless) → cor de cada quadrado exibido; painel de detalhe de uma célula (status, SD, Login, Gravação por câmera); '
         'quantidade de bolinhas; resultado mostrado no painel de manutenção; texto integral da observação do técnico (comparado caractere a caractere com a planilha).',
         '5. **Formulários**: quantidade na planilha × `manutencoes_tratadas.csv` × JSON.', '',
         f'## Resultado: {ok} de {len(res)} prefixos sem nenhuma divergência', '',
         '| Prefixo | Manutenção? | Dias conferidos (planilha = processado = JSON = painel) | Registros na planilha | Formulários (planilha/processado/JSON) | Bolinhas no painel | Resultado exibido | Observações do técnico | Divergências |',
         '|---|---|---|---|---|---|---|---|---|']
    for r in res:
        L.append(f"| {r['prefixo']} | {'sim' if r['forms_planilha'] else 'não'} | {r['dias']} | {r['registros_brutos']} | {r['forms_planilha']}/{r['forms_processados']}/{r['forms_json']} | {r['bolinhas']} | "
                 f"{', '.join(r['resultados']) or '—'} | {r['obs'] if r['forms_planilha'] else '—'} | {len(r['problemas']) or 'nenhuma'} |")
    L += ['', '## Detalhes', '']
    for r in res:
        L.append(f"- **{r['prefixo']}** — cores (01/09→25/09, V=verde L=laranja R=vermelho espaço=sem dados): `{r['cores']}`; célula conferida no painel: {r.get('celula_conferida', '—')}.")
        for pr in r['problemas']:
            L.append(f'  - ⚠️ {pr}')
    L += ['', '## Conclusão', '',
          'Os dados exibidos no painel correspondem às planilhas originais em todos os pontos conferidos acima. '
          'Observações longas aparecem completas (sem corte) no painel de manutenção.' if ok == len(res) else
          'Há divergências listadas acima que precisam ser analisadas.', '']
    (RAIZ / 'docs' / 'VALIDACAO.md').write_text('\n'.join(L), encoding='utf-8')
    print('\n'.join(L))


if __name__ == '__main__':
    main()
