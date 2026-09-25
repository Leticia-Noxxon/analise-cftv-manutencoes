"""F7 — Processamento do formulário de manutenção (Revisão_CFTV*.xlsx).

Lê TODA a aba de respostas (todas as colunas/linhas, texto integral com quebras de linha).
Colunas identificadas pelo NOME (a ordem no arquivo é irregular). Nada é truncado ou descartado:
cada formulário guarda também a lista completa de respostas originais (coluna -> valor).
"""
import datetime as dt
import re
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl

from . import config
from .interpretacao import normalizar_posicao, vazio

MESES = {'janeiro': 1, 'fevereiro': 2, 'março': 3, 'marco': 3, 'abril': 4, 'maio': 5, 'junho': 6, 'julho': 7,
         'agosto': 8, 'setembro': 9, 'outubro': 10, 'novembro': 11, 'dezembro': 12}
_RX_DATA = re.compile(r'^\s*(?:[^,]+),\s*([A-Za-zçÇ]+)\s+(\d{1,2}),\s*(\d{4})\s+(\d{1,2}):(\d{2})')
_RX_COL_POS = re.compile(r'^(Problemas Detectados|Ações)\s*:\s*C[âa]m\.?\s*(.+)$', re.IGNORECASE)
SEM_PROBLEMA = 'nenhuma anomalia identificada'
SEM_ACAO = 'nenhuma ação realizada'

# Q16 — termos de pendência (o trecho encontrado é sempre exibido)
PADROES_PENDENCIA = [
    r'\bpendente\b', r'\bpend[êe]ncias?\b', r'\baguardand\w*', r'n[ãa]o foi poss[íi]vel', r'\bfalta\b', r'\bfaltando\b',
    r'\bnecess[áa]ri[oa]\b', r'\bnecessit\w*',
    r'\bprecis\w*\s+(?:de\s+)?(?:(?:um|uma)\s+)?(?:nov[oa]|troc\w*|substitu\w*|reposi\w*|cart\w*|c[âa]mera)',
]
ACOES_PENDENCIA = ['câmera encaminhada para manutenção']

_RX_CAM = re.compile(r'\bc[aâ]m(?:[eêa]ras?|aras?)?\b\.?[\s,:]*((?:2[1-6])(?:\s*(?:,|;|e|&|/|\s)\s*2[1-6])*)\b',
                     re.IGNORECASE)
_RX_C = re.compile(r'\bC([1-6])\b')
_RX_HEX12 = re.compile(r'\b[0-9A-F]{12}\b')
_RX_SERIE = re.compile(r'\b(?=[0-9A-Z]*\d)(?=[0-9A-Z]*[A-Z])[0-9A-Z]{8,}\b')


def _sa(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')


def arquivo_manutencao(raw: Path = config.RAW):
    fs = sorted(p for p in raw.glob('*.xlsx') if unicodedata.normalize('NFC', p.name).startswith('Revisão_CFTV'))
    if not fs:
        raise FileNotFoundError('Arquivo Revisão_CFTV*.xlsx não encontrado em data/raw')
    return fs[-1]  # o mais recente pelo nome (data/hora no nome)


def parse_data(txt):
    m = _RX_DATA.match(str(txt or ''))
    if not m:
        return None
    mes = MESES.get(m.group(1).lower())
    if not mes:
        return None
    return dt.datetime(int(m.group(3)), mes, int(m.group(2)), int(m.group(4)), int(m.group(5)))


def itens(valor):
    if vazio(valor):
        return []
    return [x.strip() for x in str(valor).split('\n') if x.strip()]


def cameras_no_texto(txt):
    """Números de câmera (21–26) citados no texto livre: 'Camera 21', 'Câmeras 22 e 23', 'Câmara 23'..."""
    if vazio(txt):
        return []
    nums = []
    for g in _RX_CAM.findall(str(txt)):
        for n in re.findall(r'2[1-6]', g):
            if int(n) not in nums:
                nums.append(int(n))
    return sorted(nums)


def notacao_c(txt):
    if vazio(txt):
        return []
    return sorted(set('C' + n for n in _RX_C.findall(str(txt))))


def pendencias(txt, acoes_todas):
    trechos = []
    if not vazio(txt):
        linhas = [l.strip() for l in str(txt).split('\n') if l.strip()]
        for l in linhas:
            for p in PADROES_PENDENCIA:
                m = re.search(p, l, re.IGNORECASE)
                if m:
                    trechos.append({'trecho': l, 'termo': m.group(0), 'origem': 'Observações'})
                    break
    for a in acoes_todas:
        if a['acao'].lower() in ACOES_PENDENCIA:
            trechos.append({'trecho': f"{a['posicao']}: {a['acao']}", 'termo': a['acao'], 'origem': 'Ações'})
    return trechos


def equipamentos(txt):
    """Extrai nº de série e MAC com contexto 'instalado'/'retirado'. Texto original sempre preservado à parte."""
    res = {'instalado': {'series': [], 'macs': []}, 'retirado': {'series': [], 'macs': []},
           'nao_especificado': {'series': [], 'macs': []}}
    if vazio(txt):
        return res
    estado = 'nao_especificado'
    frag = None  # fragmento de série quebrado em duas linhas
    for linha in str(txt).split('\n'):
        low = _sa(linha).lower()
        ctx = estado
        if 'retirad' in low:
            ctx = estado = 'retirado'
        elif 'instalad' in low:
            ctx = estado = 'instalado'
        up = linha.upper()
        if 'MAC' in up:
            for mac in _RX_HEX12.findall(up):
                if mac not in res[ctx]['macs']:
                    res[ctx]['macs'].append(mac)
            frag = None
            continue
        toks = _RX_SERIE.findall(up)
        if len(toks) == 1 and toks[0] == up.strip() and len(toks[0]) <= 12:
            if frag:
                res[ctx]['series'].append(frag + toks[0])
                frag = None
            else:
                frag = toks[0]
            continue
        if frag:
            res[ctx]['series'].append(frag)
            frag = None
        for t in toks:
            if t not in res[ctx]['series']:
                res[ctx]['series'].append(t)
    if frag:
        res[estado]['series'].append(frag)
    return res


def normalizar_tecnicos(nomes):
    """Q13: unifica apenas nomes iguais ignorando maiúsculas/minúsculas (grafia mais frequente)."""
    cont = Counter(nomes)
    grupos = {}
    for n in cont:
        grupos.setdefault(n.strip().casefold(), []).append(n)
    mapa = {}
    for vs in grupos.values():
        pref = max(vs, key=lambda v: (cont[v], v))
        if len(vs) > 1:
            # prefere a grafia com iniciais maiúsculas quando empatar/for minoria
            titulo = [v for v in vs if v.strip() == v.strip().title()]
            pref = titulo[0] if titulo else pref
        for v in vs:
            mapa[v] = pref.strip()
    return mapa


def ler_formulario(path: Path = None):
    path = path or arquivo_manutencao()
    wb = openpyxl.load_workbook(path, data_only=True)
    info = {'arquivo': path.name, 'abas': wb.sheetnames, 'abas_ignoradas': []}
    aba = None
    for ws in wb.worksheets:
        cab = [str(c).strip() if c is not None else None for c in next(ws.iter_rows(max_row=1, values_only=True))]
        if 'Prefixo' in cab and any(c and _RX_COL_POS.match(c) for c in cab):
            if aba is not None:
                raise ValueError('Mais de uma aba de respostas no formulário')
            aba = ws
        else:
            info['abas_ignoradas'].append({'aba': ws.title, 'motivo': 'auxiliar (sem cabeçalho de respostas do formulário: Prefixo + Problemas/Ações por câmera)'})
    info['aba_utilizada'] = aba.title
    linhas = list(aba.iter_rows(values_only=True))
    cab = [str(c).strip() if c is not None else None for c in linhas[0]]
    info['colunas'] = [c for c in cab if c]
    forms = []
    for i, r in enumerate(linhas[1:], start=2):
        if all(vazio(v) for v in r):
            continue
        forms.append((i, {c: v for c, v in zip(cab, r) if c}))
    info['registros'] = len(forms)
    return forms, info


def processar(path: Path = None, prefixos_cftv=None):
    forms, info = ler_formulario(path)
    mapa_tec = normalizar_tecnicos([str(f.get('Nome') or '').strip() for _, f in forms])
    ids = Counter(f.get('ID') for _, f in forms if not vazio(f.get('ID')))
    prefixos_cftv = set(prefixos_cftv or [])
    saida = []
    for linha, f in forms:
        dh = parse_data(f.get('Data'))
        pos = {}
        for col, val in f.items():
            m = _RX_COL_POS.match(col)
            if not m:
                continue
            tipo = 'problemas' if m.group(1).lower().startswith('problema') else 'acoes'
            p = normalizar_posicao(m.group(2))
            if p is None:
                continue
            d = pos.setdefault(p, {'posicao': p, 'camera': config.CAMERA_DA_POSICAO[p], 'problemas': [], 'acoes': [],
                                   'colunas_origem': []})
            d['colunas_origem'].append(col)
            for it in itens(val):
                if it not in d[tipo]:
                    d[tipo].append(it)
        posicoes = []
        for p in config.ORDEM_POSICOES:
            if p in pos and (pos[p]['problemas'] or pos[p]['acoes']):
                d = pos[p]
                d['anomalias'] = [x for x in d['problemas'] if x.lower() != SEM_PROBLEMA]
                d['acoes_realizadas'] = [x for x in d['acoes'] if x.lower() != SEM_ACAO]
                d['tem_anomalia'] = bool(d['anomalias'])
                d['rotulo'] = config.rotulo_camera(d['camera'])
                posicoes.append(d)
        obs = None if vazio(f.get('Observações')) else str(f.get('Observações'))
        acoes_todas = [{'posicao': d['rotulo'], 'acao': a} for d in posicoes for a in d['acoes_realizadas']]
        cams_txt = cameras_no_texto(obs)
        pend = pendencias(obs, acoes_todas)
        equip = equipamentos(obs)
        pecas = []
        for d in posicoes:
            for a in d['acoes_realizadas']:
                if re.search(r'substitui|inser|instala', a, re.IGNORECASE):
                    pecas.append(f"{a} — {d['rotulo']}")
        qtd = {k: (None if vazio(f.get(k)) else str(f.get(k))) for k in
               ['Câmera Utilizada', 'Switch Utilizado', 'Cartão de Memória Utilizado']}
        prefixo = int(f['Prefixo'])
        alertas = []
        nome_orig = str(f.get('Nome') or '').strip()
        if mapa_tec.get(nome_orig, nome_orig) != nome_orig:
            alertas.append(f"Nome do técnico padronizado: '{nome_orig}' → '{mapa_tec[nome_orig]}' (diferença só de maiúsculas)")
        if dh is None:
            alertas.append('Data/hora não interpretada — verificar texto original')
        idv = f.get('ID')
        if not vazio(idv) and ids[idv] >= 10:
            alertas.append(f'ID {idv} repetido em {ids[idv]} formulários (provável valor genérico)')
        if obs:
            m = re.match(r'^\s*(\d{5})\s*-\s*(\d{5})\b', obs)
            if m and int(m.group(1)) == prefixo and not vazio(idv) and int(m.group(2)) != int(idv):
                alertas.append(f'Texto informa ID {m.group(2)}, diferente da coluna ID ({idv})')
            for n in re.findall(r'(?<![\dA-Z])(\d{5})(?![\dA-Z])', obs):
                n = int(n)
                if n != prefixo and (vazio(idv) or n != int(idv)) and n in prefixos_cftv:
                    alertas.append(f'Texto cita o prefixo {n}, diferente da coluna Prefixo ({prefixo}); mantido o valor da coluna Prefixo (Q11)')
            if re.search(r'prefixo incorreto', obs, re.IGNORECASE):
                alertas.append('Texto contém "PREFIXO INCORRETO" — conferir')
        if prefixos_cftv and prefixo not in prefixos_cftv:
            alertas.append('Prefixo não existe em nenhum Relatório CFTV')
        saida.append({
            'id_form': linha,  # linha no Excel (rastreabilidade)
            'linha_excel': linha,
            'arquivo_origem': info['arquivo'],
            'prefixo': prefixo,
            'tecnico': mapa_tec.get(nome_orig, nome_orig),
            'tecnico_original': nome_orig,
            'data_texto': f.get('Data'),
            'datahora': dh.strftime('%Y-%m-%d %H:%M') if dh else None,
            'data': dh.strftime('%Y-%m-%d') if dh else None,
            'hora': dh.strftime('%H:%M') if dh else None,
            'id_formulario': None if vazio(idv) else idv,
            'garagem': None if vazio(f.get('Garagem')) else str(f.get('Garagem')).strip(),
            'tecnologia': None if vazio(f.get('Tecnologia')) else str(f.get('Tecnologia')).strip(),
            'posicoes': posicoes,
            'cameras_com_anomalia': [d['camera'] for d in posicoes if d['tem_anomalia']],
            'cameras_citadas_texto': cams_txt,
            'notacao_c_texto': notacao_c(obs),
            'quantidades': qtd,
            'pecas': pecas,
            'equipamento': equip,
            'pendencia': bool(pend),
            'pendencia_trechos': pend,
            'observacoes': obs,
            'respostas': [{'coluna': c, 'valor': (v.strftime('%Y-%m-%d %H:%M') if isinstance(v, dt.datetime) else v)}
                          for c, v in f.items() if not vazio(v)],
            'alertas': alertas,
        })
    info['tecnicos_padronizados'] = {k: v for k, v in mapa_tec.items() if k != v}
    return saida, info
